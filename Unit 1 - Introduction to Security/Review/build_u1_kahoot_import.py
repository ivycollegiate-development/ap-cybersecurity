#!/usr/bin/env python3
"""Build Kahoot! import spreadsheet for AP Cyber Unit 1 review (test Mon Sep 28).

Source material:
- Unit 1 Test Bank MCQs (Unit-1-Lesson-Plans.md, U1 L18 section) — questions 1-5
- U1 L16 Review Worksheet quick checks (decks 1.1-1.5) — questions 6-11

Kahoot import constraints (kahoot.com importer):
- .xlsx, question <= 95 chars, each answer <= 60 chars
- >= 2 answers, correct answer(s) column = number 1-4
- Template columns: Question | Answer 1-4 | Time limit in seconds | Correct answer(s)
"""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

OUT = ("/Users/pauljonessr/Projects/ap-cybersecurity/"
       "Unit 1 - Introduction to Security/Review/U1-Kahoot-Import.xlsx")

Q = [
    # ---- Unit 1 Test Bank (lesson plans U1 L18) ----
    ("Which of the following is the strongest indicator of a phishing email?",
     ["The email uses the company logo",
      "The address contains a subtle misspelling (g00gle.com)",
      "The email is addressed to you by name",
      "The email contains an unsubscribe link"], 2, "1.1"),

    ("A user notices logins from unfamiliar IP addresses during school hours. Most likely a sign of:",
     ["A brute force attack in progress",
      "Successful credential compromise",
      "Normal account activity",
      "A DDoS attack"], 2, "1.1"),

    ("Which of the following is NOT a recommended practice on public Wi-Fi?",
     ["Use a VPN",
      "Connect to any open network with \u201cFree Wi-Fi\u201d in the name",
      "Ensure websites use HTTPS",
      "Turn off file sharing"], 2, "1.3"),

    ("A hotspot broadcasts the same network name as a coffee shop's Wi-Fi. This attack is:",
     ["A brute force attack",
      "An evil twin attack",
      "A dictionary attack",
      "A credential stuffing attack"], 2, "1.3"),

    ("Which adversary type is most motivated by financial gain?",
     ["Hacktivists",
      "Script kiddies",
      "Transnational criminal organizations",
      "Cyberterrorists"], 3, "1.1"),

    # ---- U1 L16 Review Worksheet quick checks ----
    ("Which of the following is the clearest indicator of a social engineering attack?",
     ["An email that arrives during school hours",
      "Urgency that pressures you to act before checking",
      "A newsletter you subscribed to",
      "An email with an attachment you requested"], 2, "1.1"),

    ("Maya's cat's name, jersey number, and school are all public. What protects her accounts?",
     ["Posting less often so attackers find fewer facts",
      "Building passwords from non-public information",
      "Adding more friends she knows in real life",
      "Changing her username every month"], 2, "1.2"),

    ("Why can a single data breach lead to multiple accounts being compromised?",
     ["Because strong passwords are easy for attackers to guess",
      "Because attackers test it on many other websites",
      "Because users reuse one password on many accounts",
      "Because all accounts are connected automatically"], 3, "1.2"),

    ("A website you visit on cafe Wi-Fi shows a certificate warning. What should you do?",
     ["Click through it - warnings are usually fake",
      "Leave the site; the connection may be intercepted or spoofed",
      "Reload the page until the warning disappears",
      "Enter your password quickly before the warning returns"], 2, "1.3"),

    ("Which of these is an example of adversaries using AI to augment an ATTACK?",
     ["An AI tool flagging SQL injection in code before release",
      "An AI writing phishing emails with zero grammar errors",
      "An SOC dashboard alerting on unusual process behavior",
      "A password manager generating a long random password"], 2, "1.4"),

    ("AI makes AI-capable tasks free. What happens to human-accountable tasks' value?",
     ["It decreases, because AI eventually does everything",
      "It stays exactly the same",
      "It increases - human judgment becomes MORE valuable",
      "It disappears, because AI handles accountability too"], 3, "1.5"),

    # ---- Added for balance: flagged hardest LOs (1.4 attack side, 1.5 defense side) ----
    ("Deepfake voice scams (like the fake-parent call) work because AI can:",
     ["Clone a voice from samples posted online",
      "Break passwords instantly",
      "Block scam calls automatically",
      "Translate voices in real time"], 1, "1.4"),

    ("Which of these is a DEFENSIVE use of AI?",
     ["Flagging SQL injection in code before release",
      "Cloning a CEO's voice to approve transfers",
      "Writing phishing emails at scale",
      "Scraping profiles to guess passwords"], 1, "1.5"),

    ("OpenAI's postmortem: agents evaded monitoring for weeks. The defense lesson:",
     ["Monitoring only works if alerts reach humans fast",
      "Monitoring is unnecessary for AI systems",
      "Agents can never disable their own monitors",
      "Logs alone are enough; audits are optional"], 1, "1.5"),
]

# --- enforce Kahoot limits BEFORE writing anything ---
for q, answers, correct, tag in Q:
    assert len(q) <= 95, f"question too long ({len(q)}): {q}"
    for a in answers:
        assert len(a) <= 60, f"answer too long ({len(a)}): {a}"
    assert len(answers) == 4 and 1 <= correct <= 4

wb = Workbook()
ws = wb.active
ws.title = "Sheet1"
ws.append(["Question", "Answer 1", "Answer 2", "Answer 3", "Answer 4",
           "Time limit in seconds", "Correct answer(s)"])
for q, answers, correct, tag in Q:
    ws.append([q] + list(answers) + [20, correct])  # 20s default per Kahoot

hdr_fill = PatternFill("solid", fgColor="1F3864")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = hdr_fill
ws.freeze_panes = "A2"
ws.column_dimensions["A"].width = 55
for col in "BCDE":
    ws.column_dimensions[col].width = 42
ws.column_dimensions["F"].width = 12
ws.column_dimensions["G"].width = 12
# teacher-only column: CED topic tag (extra col; importer ignores extras per template docs)
ws.cell(row=1, column=8, value="CED topic (not imported)")
for i, (q, a, c, t) in enumerate(Q, start=2):
    ws.cell(row=i, column=8, value=t)

wb.save(OUT)
print("saved:", OUT)

# --- read back and verify ---
ws2 = load_workbook(OUT).active
rows = list(ws2.iter_rows(values_only=True))
assert rows[0][:7] == ("Question", "Answer 1", "Answer 2", "Answer 3", "Answer 4",
                       "Time limit in seconds", "Correct answer(s)")
probs = []
for r in rows[1:]:
    if len(r[0]) > 95:
        probs.append(("question", r[0]))
    if any(len(str(x)) > 60 for x in r[1:5]):
        probs.append(("answer", r[0]))
    if r[5] != 20 or r[6] not in (1, 2, 3, 4):
        probs.append(("meta", r[0]))
print("questions:", len(rows) - 1)
print("constraint check:", "ALL PASS" if not probs else probs)
from collections import Counter
print("topic spread:", dict(Counter(r[7] for r in rows[1:])))
